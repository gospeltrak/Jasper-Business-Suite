from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/msamimsechu/Jasper-Business-Suite/Jasper-Business-Suite')
SHOTS = ROOT / 'artifacts/orvix-guide-detailed/screenshots'
OUT = ROOT / 'output/docx/Orvix_Mwongozo_Kamili_wa_Mobile.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY='11182D'; GREEN='00A878'; VIOLET='7652FF'; BLUE='2E74B5'; MUTED='66758F'; PALE='EAFBF5'; LIGHT='F4F7FB'; WHITE='FFFFFF'; AMBER='FFF5D6'; RED='FDEBEC'

doc=Document(); sec=doc.sections[0]
sec.page_height=Inches(11); sec.page_width=Inches(8.5)
sec.top_margin=sec.bottom_margin=Inches(1); sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(11); styles['Normal'].font.color.rgb=RGBColor.from_string(NAVY)
styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.25
for name,size,color,before,after in [('Title',30,NAVY,0,12),('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',12,'1F4D78',10,5)]:
    s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)

def shade(cell,color):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(shd)
def margins(cell,top=80,bottom=80,start=120,end=120):
    pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); pr.append(mar)
    for k,v in [('top',top),('bottom',bottom),('start',start),('end',end)]:
        n=mar.find(qn('w:'+k))
        if n is None: n=OxmlElement('w:'+k); mar.append(n)
        n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa')
def bookmark(p,name,i):
    a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(i)); a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'),str(i)); p._p.insert(0,a); p._p.append(b)
def link(p,text,anchor):
    h=OxmlElement('w:hyperlink'); h.set(qn('w:anchor'),anchor); h.set(qn('w:history'),'1')
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),GREEN); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.extend([c,u])
    t=OxmlElement('w:t'); t.text=text; r.extend([rp,t]); h.append(r); p._p.append(h)
def box(title,text,color=PALE):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.5)
    c=t.cell(0,0); shade(c,color); margins(c); p=c.paragraphs[0]; rr=p.add_run(title.upper()+'  '); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(GREEN if color!=RED else 'B42318'); p.add_run(text)
def steps(items):
    for x in items:
        p=doc.add_paragraph(style='List Number'); p.add_run(x)
def bullets(items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.add_run(x)
def shot(file,caption,width=3.0):
    path=SHOTS/file
    if not path.exists(): return
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
    p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(8.5); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
def chapter(num,title,intro,image=None,caption=None):
    doc.add_page_break(); h=doc.add_heading(f'{num}. {title}',1); bookmark(h,f's{num}',100+num)
    p=doc.add_paragraph(intro); p.runs[0].font.size=Pt(11.5); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    if image: shot(image,caption or title)
def action(title,goal,items,result=None,warn=None):
    doc.add_heading(title,2); doc.add_paragraph(goal); steps(items)
    if result: box('Matokeo',result)
    if warn: box('Tahadhari',warn,RED)

# Header/footer
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=h.add_run('ORVIX  •  MWONGOZO KAMILI WA MOBILE'); r.bold=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GREEN)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('orvix.africa  •  info@orvix.africa  •  +255 655 746 552'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

# Cover — editorial cover pattern
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ROOT/'public/icon-512.png'),width=Inches(1.25))
p=doc.add_paragraph('ORVIX',style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('MWONGOZO KAMILI WA MTUMIAJI WA MOBILE'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.runs[0]; r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(GREEN)
p=doc.add_paragraph('Screen kwa screen • Menu kwa menu • Hatua kwa hatua'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
doc.add_paragraph('\n'); box('Toleo la mafunzo — August 2026','Limetengenezwa kwa mobile app ya sasa na Google Sign-In.')
doc.add_paragraph('\n')
p=doc.add_paragraph('FARAGHA YA DATA'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True; p.runs[0].font.color.rgb=RGBColor.from_string(VIOLET)
p=doc.add_paragraph('Majina, barua pepe, simu, kampuni, wateja na miamala yote kwenye picha na mifano ni ya kubuni kwa ajili ya demo. Hakuna taarifa ya mtumiaji halisi iliyotumika.'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break(); doc.add_heading('Jinsi ya kutumia kitabu hiki',1)
bullets(['Fuata namba za hatua kwa mpangilio.','Maneno yenye rangi ya kijani kwenye Yaliyomo yanakupeleka moja kwa moja kwenye sehemu husika.','Picha zinaonyesha screen halisi ya mobile; data ni ya demo.','Kama button haionekani, role yako inaweza kutokuwa na ruhusa hiyo. Wasiliana na Admin.'])
box('Alama zinazotumika','Matokeo = unachotakiwa kuona baada ya hatua. Tahadhari = hatua yenye athari kama kufuta, kurejesha stock au kubadilisha malipo.',AMBER)

toc=[(1,'Kusajili na kuingia kwa Google'),(2,'Home / Dashboard'),(3,'Sell / POS'),(4,'Checkout na malipo'),(5,'Sales na utafutaji'),(6,'Kufungua sale na vitendo vyake'),(7,'Stock na bidhaa'),(8,'Buying / Purchases'),(9,'Expenses'),(10,'Reports'),(11,'Partners'),(12,'Delivery'),(13,'Money'),(14,'Planning'),(15,'Staff, role na invite'),(16,'Settings, Sync, Subscription na App'),(17,'Utatuzi wa matatizo'),(18,'Mawasiliano')]
doc.add_page_break(); doc.add_heading('Yaliyomo',1)
for n,tit in toc:
    t=doc.add_table(rows=1,cols=2); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.columns[0].width=Inches(.55); t.columns[1].width=Inches(5.95)
    shade(t.cell(0,0),GREEN); shade(t.cell(0,1),'E8EEF5'); margins(t.cell(0,0)); margins(t.cell(0,1))
    p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run(str(n)); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(WHITE)
    link(t.cell(0,1).paragraphs[0],tit,f's{n}')

chapter(1,'Kusajili na kuingia kwa Google','Orvix ya sasa haitumii namba ya simu na password kama njia ya kuingia. Account huunganishwa na Google.', '02-create-account-google.png','Create Account: anza kwa Continue with Google. Data yoyote inayoonekana ni ya demo.')
action('1.1 Kutengeneza account mpya','Tumia Gmail unayotaka iwe utambulisho wa Orvix.',[
    'Fungua link ya Orvix kwenye Chrome au Safari ya simu.','Bonyeza Create Account.','Bonyeza Continue with Google.','Chagua Google account yako kwenye dirisha la Google.','Baada ya kurudi Orvix, jaza Business Name, Business Type, City/Office na simu ya biashara ikiwa form inaomba.','Soma na ukubali masharti, kisha bonyeza Create Account.'], 'Dashboard itafunguka na tenant mpya pamoja na profile ya Admin.','Usitumie Gmail ya mtu mwingine. Orvix haiombi Google password ndani ya form yake.')
action('1.2 Kuingia baada ya kusajili','Ingia kwenye kifaa chochote kwa Google account ileile.',[
    'Fungua Orvix na bonyeza Sign In.','Bonyeza Continue with Google.','Chagua Gmail ileile iliyotumika kusajili au kukubali invite.','Subiri logo ya tenant wakati menus, role na data vinafunguliwa.'], 'Unaingia kwenye tenant na branch unazoruhusiwa. Hakuna ujumbe wa “restore menus/roles” unaohitajika kuonekana.')
action('1.3 Staff anapotumia invite','Staff hatengenezi tenant mpya.',[
    'Admin anatengeneza invite kwenye Staff.','Staff anafungua link hiyo.','Staff anachagua Continue with Google na Gmail iliyolengwa na invite.','Baada ya kukamilisha, anaingia kwenye profile yenye role na branch alizopewa.'], 'Profile ya staff inabaki database na inaweza kutumika kwenye kifaa tofauti.')

chapter(2,'Home / Dashboard','Home ni muhtasari wa biashara na njia ya haraka ya kufungua POS.','03-dashboard-home.png','Dashboard ya demo kwenye simu.')
action('2.1 Kusoma summary cards','Elewa hali ya biashara kwa kipindi kilichochaguliwa.',[
    'Angalia Total Orders, Total Sales, Purchases, Expenses, Dues Owed na Total Profit.','Badilisha Today, 1 Week, 1 Month, 3 Month au 1 Year.','Hakiki branch iliyo juu kabla ya kutafsiri takwimu.'])
action('2.2 Kufungua screen ya kuuza','Anza sale haraka.', ['Bonyeza Open Sell Screen/Start Selling.','Utaingia POS yenye search, product grid na cart.'])
action('2.3 Bottom navigation','Tumia tabs tano za kudumu.', ['Home: dashboard.','Sales: historia, receipt, debt na quote.','Sell: POS.','Stock: bidhaa na quantity.','More: menu nyingine zote.'])

chapter(3,'Sell / POS','POS ndiyo sehemu ya kutafuta bidhaa na kuziweka kwenye cart.','05-pos-products.png','Product grid ya POS ya demo.')
action('3.1 Kutafuta bidhaa','Tafuta kwa jina, code au barcode.',[
    'Bonyeza Sell.','Gusa Search Code, Barcode or Title.','Andika sehemu ya jina, code au barcode.','Gusa bidhaa sahihi kwenye matokeo.','Futa maandishi ya search kurudisha bidhaa zote.'], 'Matokeo hupungua kadri unavyoandika; screen haitakiwi kuwa nyeupe kwenye iPhone, iPad au Android.')
action('3.2 Scanner','Soma barcode kwa camera.',[
    'Bonyeza icon ya Scan.','Ruhusu camera ikiwa browser inaomba.','Elekeza camera kwenye barcode.','Thibitisha bidhaa iliyopatikana kabla ya kuiweka cart.'])
action('3.3 Retail na Wholesale','Chagua aina ya bei kabla ya kuongeza bidhaa.',[
    'Bonyeza Retail Sales type kwa mauzo ya kawaida.','Au bonyeza Wholesale Sales type kwa bei ya jumla.','Hakiki badge ya aina ya sale inayoonekana.'])
action('3.4 Kuongeza na kubadilisha cart','Dhibiti bidhaa kabla ya checkout.',[
    'Bonyeza + Add kwenye product card.','Bonyeza + au − kubadilisha quantity.','Tumia remove/trash kuondoa line item.','Weka discount kwa asilimia au cash ikiwa umepewa ruhusa.','Chagua VAT/No Tax na delivery charge kama vinahusika.'], 'Cart Total na Order Total husasishwa papo hapo.')
shot('06-pos-cart-item.png','Bidhaa ya demo imeongezwa kwenye cart.')

chapter(4,'Checkout na malipo','Checkout hukusanya taarifa za mteja, njia ya malipo na uthibitisho wa sale.','07-checkout-payment-methods.png','Hatua ya kuchagua payment method.')
action('4.1 Kufungua checkout','Hakiki cart kwanza.', ['Thibitisha bidhaa, quantity, discount, tax na delivery.','Bonyeza Checkout.','Chagua au andika customer kwa kutumia data sahihi ya biashara.'])
action('4.2 Kuchagua njia ya malipo','Rekodi njia halisi iliyotumika.',[
    'Cash: mteja amelipa fedha taslimu.','Card: malipo ya kadi.','MOMO/Mobile Money: malipo ya simu.','Credit: kiasi kinabaki deni.','Weka amount received au due kulingana na screen, kisha hakiki balance/change.'])
shot('08-checkout-filled-demo.png','Checkout iliyojazwa kwa data ya kubuni: Demo Customer.')
action('4.3 Kukamilisha sale','Hii ndiyo hatua inayotengeneza receipt na kupunguza stock.',[
    'Soma total mara ya mwisho.','Bonyeza Complete/Confirm Payment mara moja.','Subiri TRANSACTION APPROVED.','Tuma PDF kupitia WhatsApp ikiwa mteja ameomba, au bonyeza Finish Sale.'], 'Receipt mpya ina invoice number, cashier, customer, bidhaa, total na payment details.','Usibonyeze confirm mara nyingi. Ukiona tatizo, hakiki Sales kabla ya kurudia.')
shot('09-sale-receipt-success.png','Receipt ya muamala wa demo. Majina na mawasiliano ya mfano si ya mtumiaji halisi.')

chapter(5,'Sales na utafutaji','Sales inatunza historia ya miamala na shortcuts za receipt, debt, settlement na quote.','10-sales-with-transaction.png','Sales Overview yenye muamala mmoja wa demo.')
action('5.1 Kutafuta sale','Tafuta bila kuscroll list ndefu.',[
    'Bonyeza Sales.','Gusa Search sales.','Andika customer, invoice/reference au sehemu ya maelezo.','Tumia filter ya Cash, Card, MOMO au Credit.','Tumia Today, Week, Month au All.'])
action('5.2 Receipts','Fungua orodha ya risiti.', ['Bonyeza Receipts.','Chagua receipt.','Tumia POS Receipt kwa thermal slip au A4 Invoice kwa invoice kubwa.','Print, download au share kulingana na kifaa.'])
action('5.3 Debts na Settle','Fuatilia credit sales na malipo ya deni.',[
    'Bonyeza Debts kuona balances.','Tafuta customer au invoice.','Bonyeza Settle/Add Payment.','Weka kiasi na payment method.','Thibitisha balance mpya kabla ya kuhifadhi.'], warn='Usirekodi settlement bila uthibitisho wa malipo.')
action('5.4 Quotes','Tengeneza na kufuatilia quotation.', ['Bonyeza Quotes.','Fungua quote iliyopo au New Quote.','Ongeza customer, bidhaa, quantity, bei na validity.','Save, share au convert to sale inapokubaliwa.'])

chapter(6,'Kufungua sale na vitendo vyake','Kila sale ina action menu ya kuona, kulipa, kuhariri, kuchapisha, delivery na cancel.','11-sale-actions.png','Action menu ya muamala wa demo.')
action('6.1 View Sale','Ona rekodi bila kuibadilisha.', ['Kwenye Sales, bonyeza Sale actions.','Bonyeza View Sale.','Soma Date & Time, Payment Method, Cashier, Customer na Items Purchased.','Linganisha Sub-total, Grand Total, Paid na Balance Due.'])
shot('12-sale-detail-products.png','View Sale inaonyesha bidhaa, quantity, bei, paid na balance.')
action('6.2 Edit Sale','Sahihisha customer au items ikiwa role inaruhusu.', ['Fungua Sale actions.','Bonyeza Edit Sale.','Badilisha field inayohitajika tu.','Hakiki total na stock impact.','Save changes.'], warn='Edit inaweza kubadilisha stock na accounting. Tumia kwa correction halisi tu.')
action('6.3 POS Receipt na A4 Invoice','Chagua format ya kuchapisha.', ['POS Receipt: receipt nyembamba ya printer.','A4 Invoice: ukurasa kamili wa invoice.','Tumia Share/Print/Download ya browser au simu.'])
action('6.4 Add to Delivery','Peleka sale kwenye dispatch.', ['Bonyeza Add to Delivery.','Jaza recipient, simu ya demo/halisi ya mteja, address na delivery note.','Save.','Fuatilia status kwenye Delivery.'])
action('6.5 Cancel Receipt','Reverse sale salama.', ['Bonyeza Cancel Receipt.','Soma kiasi na stock zitakazorejeshwa.','Andika reason.','Thibitisha mara moja.'], warn='Cancel Receipt ni action kubwa. Usiitumie kufuta historia au kuficha makosa.')

chapter(7,'Stock na bidhaa','Stock ni catalogue, bei, quantity, category, barcode na movement za bidhaa.','03-dashboard-home.png','Tumia Stock kwenye bottom navigation; picha hii inaonyesha eneo la nav ya mobile.')
action('7.1 Kutafuta bidhaa','Bonyeza Stock, gusa search, kisha andika jina, SKU au barcode. Tumia category/status filters kupunguza matokeo.', ['Fungua Stock.','Andika search term.','Fungua product card sahihi.'])
action('7.2 Kuongeza bidhaa','Jaza taarifa zote za msingi.',[
    'Bonyeza Add/Register New Product.','Weka Product Name, SKU/Barcode na Category.','Weka Unit, Buying Price, Retail Price na Wholesale Price.','Weka Opening Stock, Low-stock level na branch.','Ongeza picha ikiwa inahitajika.','Bonyeza Save na utafute bidhaa kuthibitisha.'], warn='Tumia SKU/barcode ya kipekee; duplicate inaweza kuvuruga search na scanner.')
action('7.3 Kuhariri bidhaa','Fungua bidhaa, bonyeza Edit, badilisha field, Save, kisha hakiki POS.', ['Badilisha jina/category kwa tahadhari.','Bei mpya itatumika kwa sales mpya.','Usibadilishe opening stock kusahihisha movement ya zamani.'])
action('7.4 Adjust Stock','Rekodi kuongeza au kupunguza quantity na reason.', ['Fungua bidhaa.','Bonyeza Adjust Stock.','Chagua Increase au Decrease.','Weka quantity na sababu.','Confirm na hakiki stock movement.'])
action('7.5 Kufuta bidhaa','Tumia delete ikiwa bidhaa kweli haitakiwi.', ['Fungua product actions.','Bonyeza Delete.','Soma confirmation na dependencies.','Confirm tu baada ya kuhakiki.'], warn='Bidhaa yenye sales/purchases mara nyingi ni bora ku-disable/archive kuliko kufuta.')
action('7.6 Categories, brands, barcode na import','Tumia tools za catalogue.', ['Categories/Brands: add, edit au archive grouping.','Barcode/Labels: chagua products na template, preview, kisha print.','Import: pakua template, jaza columns bila kubadilisha headers, upload na rekebisha validation errors.'])

chapter(8,'Buying / Purchases','Purchases huingiza stock kutoka supplier na kurekodi deni au malipo.')
action('8.1 Purchase mpya','Fungua More → Buying → New Purchase.', ['Chagua supplier.','Weka invoice/reference na tarehe.','Ongeza bidhaa, quantity, buying price na discount/tax.','Weka payment status: Paid, Partial au Credit.','Hakiki total na Save/Complete.'], 'Stock ya bidhaa huongezeka baada ya purchase kukamilika.')
action('8.2 Kutafuta, kufungua na kuhariri','Tumia search/filter, fungua purchase, kisha action menu.', ['View: supplier, items, total, paid na balance.','Add Payment: rekodi installment.','Edit: sahihisha taarifa zinazoruhusiwa.','Receipt/Invoice: print au share.','Cancel/Delete: tumia kwa ruhusa na reason.'])

chapter(9,'Expenses','Expenses ni matumizi ya biashara yasiyo purchase ya stock.')
action('9.1 Kuongeza expense','More → Expenses → Add Expense.', ['Chagua category.','Weka amount, date na payment account/method.','Andika description/reference.','Ongeza attachment ikiwa ipo.','Save na hakiki list.'])
action('9.2 Edit, search na delete','Tafuta kwa maelezo, category au period.', ['Fungua expense.','Edit kwa correction.','Delete/void baada ya confirmation ikiwa imeingizwa kimakosa.'], warn='Expense halisi isifutwe ili kubadilisha report; tengeneza correction yenye audit trail.')

chapter(10,'Reports','Reports hubadilisha data ya sales, purchases, stock na expenses kuwa maamuzi.')
action('10.1 Kuendesha report','More → Reports.', ['Chagua aina ya report.','Weka Today/Week/Month/Custom date.','Chagua branch, category, product au staff ikiwa filter ipo.','Bonyeza Apply/Generate.','Export/Share ikiwa inahitajika.'])
action('10.2 Reports kuu','Soma kila report kwa lengo lake.', ['Sales: revenue, quantity, payment methods na trends.','Purchases: buying totals na suppliers.','Profit: sales minus cost/expenses kulingana na calculation.','Stock: on-hand, low stock, valuation na movements.','Debts: receivables/payables na balances.','Expense: category na period.'])

chapter(11,'Partners','Partners ni Customers na Suppliers.')
action('11.1 Kuongeza partner','More → Partners → Add.', ['Chagua Customer au Supplier.','Weka jina la demo/halali, simu, email, address na tax details ikiwa zipo.','Save na utafute partner kuthibitisha.'])
action('11.2 Edit, statement na delete','Fungua partner card.', ['Edit hubadilisha mawasiliano.','Statement/history huonyesha sales, purchases, payments au balance.','Delete/archive hutolewa kwenye action menu.'], warn='Linda taarifa binafsi; usiweke data ya mteja kwenye screenshots za mafunzo.')

chapter(12,'Delivery','Delivery inasimamia orders kutoka dispatch hadi kukabidhiwa.')
action('12.1 Kutengeneza delivery','Tumia Add to Delivery kutoka sale au More → Delivery → New.', ['Chagua sale/order.','Weka recipient, phone, address, rider/vehicle na tarehe.','Save.'])
action('12.2 Kubadilisha status','Fungua delivery na chagua status.', ['Pending: haijaanza.','Dispatched/In Transit: imetoka.','Delivered: imekabidhiwa.','Failed/Returned: haikufanikiwa; andika reason.'])

chapter(13,'Money','Money ina cash, bank/mobile accounts na movements.')
action('13.1 Akaunti na transaction','More → Money.', ['Chagua account.','Tumia Add Transaction/Transfer kulingana na ruhusa.','Weka amount, source/destination, date na reference.','Save na hakiki balance.'])
action('13.2 Reconciliation','Linganisha Orvix na statement halisi.', ['Chagua period.','Linganisha entries moja moja.','Tambua missing/duplicate.','Tengeneza correction yenye note.'])

chapter(14,'Planning','Planning husaidia kuweka forecast, budget au target zinazopatikana kwenye package yako.')
action('14.1 Kutengeneza mpango','More → Planning.', ['Chagua aina ya plan.','Weka period, branch na amount/metric.','Save.','Rudi baadaye kulinganisha actual dhidi ya plan.'])

chapter(15,'Staff, role na invite','Staff access lazima ianze kwa invite, Google account, role na branch zilizoainishwa.')
action('15.1 Kutengeneza invite','More → Staff → Invite/Add Staff.', ['Weka jina la demo au la staff, na Gmail atakayotumia.','Chagua role.','Chagua branch(es).','Hakiki permissions za role.','Bonyeza Create/Send Invite.','Copy link na uitume kwa staff husika tu.'], 'Link inaunganisha usajili wa Google na tenant, role pamoja na branch.')
action('15.2 Staff kukamilisha','Staff afungue link kwenye browser.', ['Bonyeza Continue with Google.','Chagua Gmail iliyolengwa.','Kubali hatua zinazohitajika.','Ingia; logo ya tenant itaonekana wakati workspace inafunguliwa.'])
action('15.3 Roles na permissions','Admin abadilishe access kwa kiwango kinachohitajika tu.', ['Fungua staff profile.','Chagua role au Custom Permissions.','Washa/zima View, Create, Edit, Delete/Cancel kwa modules husika.','Save na mwambie staff aingie upya ikiwa access haijasasishwa.'])
action('15.4 Kifaa tofauti na profile isiyopotea','Google identity ndiyo key ya staff.', ['Tumia Gmail ileile kwenye simu/tablet/computer nyingine.','Continue with Google.','Orvix hupakia tenant, role na branch kutoka database.'], 'Profile haitegemei local browser storage pekee; haitakiwi kujifuta baada ya refresh.')

chapter(16,'Settings, Sync, Subscription na App','System tools ziko chini ya More.')
action('16.1 Settings','Badilisha taarifa za biashara kwa tahadhari.', ['More → Settings.','Fungua Business/Profile, Branch, Tax, Receipt, Language au modules zinazopatikana.','Badilisha field na Save.','Refresh na hakiki imebaki.'])
action('16.2 Sync','Angalia hali ya data.', ['More → Sync.','All data synced = mabadiliko yamefika cloud.','Ikiwa pending/offline, rudisha internet na jaribu sync tena bila kufunga app.'])
action('16.3 Subscription','Simamia package.', ['More → Subscription.','Angalia plan, expiry na amount.','Pay online au upload receipt kulingana na option.','Subiri uthibitisho na hakiki status.'])
action('16.4 Install Orvix App','Ongeza web app kwenye Home Screen.', ['Android/Chrome: menu → Add to Home screen/Install app.','iPhone/iPad Safari: Share → Add to Home Screen → Add.','Fungua icon mpya na uingie kwa Google.'])

chapter(17,'Utatuzi wa matatizo','Jaribu hatua hizi kabla ya kuwasiliana na support.')
rows=[('Search inakuwa blank','Funga keyboard, rudi screen iliyopita na ufungue search tena; refresh app; hakiki internet; sasisha Safari/iOS.'),('Error 500 baada ya reload','Subiri sekunde chache na refresh mara moja; hakiki URL/domain; tuma screenshot na saa ya tukio.'),('Menus chache','Hakiki tenant/branch; sign out na Continue with Google; Admin ahakiki role na permissions.'),('Staff haonekani','Admin ahakiki Gmail ya invite, status na tenant; staff atumie Gmail ileile.'),('Data haijasync','Acha app wazi, rudisha internet, fungua More → Sync.')]
t=doc.add_table(rows=1,cols=2); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.columns[0].width=Inches(2); t.columns[1].width=Inches(4.5)
for i,x in enumerate(['Tatizo','Hatua za kwanza']): shade(t.cell(0,i),'E8EEF5'); t.cell(0,i).text=x; t.cell(0,i).paragraphs[0].runs[0].bold=True
for a,b in rows:
    c=t.add_row().cells; c[0].text=a; c[1].text=b
    for x in c: margins(x)
box('Ujumbe wa support','Tuma screenshot, menu uliyokuwa, hatua ulizofanya, aina ya kifaa/browser na muda wa tukio. Usitume password, token au taarifa binafsi za mteja.',AMBER)

chapter(18,'Mawasiliano','Kwa msaada wa Orvix, tumia njia rasmi.')
box('Orvix Solutions','P.O. Box 1259, Mbeya, Tanzania\nWhatsApp / Simu: +255 655 746 552\nEmail: info@orvix.africa\nWebsite: https://orvix.africa/',PALE)
doc.add_heading('Checklist ya mwisho kwa mtumiaji mpya',2)
bullets(['Nimeingia kwa Google account sahihi.','Nimehakiki tenant na branch.','Ninaona menus zinazoruhusiwa na role yangu.','Nimejaribu search, sale ya mafunzo na receipt.','Ninaelewa namna ya kutafuta sale na kufungua bidhaa zake.','Najua wapi kupata Stock, Buying, Expenses, Reports na Staff.','Sitashiriki invite link, taarifa za mteja au credentials hadharani.'])

doc.save(OUT)
print(OUT)
