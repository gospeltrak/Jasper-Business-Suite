from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/msamimsechu/Jasper-Business-Suite/Jasper-Business-Suite')
SHOTS = ROOT / 'artifacts/orvix-guide/screenshots'
OUT = ROOT / 'output/docx/Orvix_Mwongozo_wa_Mobile.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)

GREEN = '00A878'; DARK = '11182D'; MUTED = '7183A1'; PALE = 'EAFBF5'; VIOLET = '7652FF'; LIGHT = 'F5F7FB'; WHITE = 'FFFFFF'

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(.55)
sec.left_margin = sec.right_margin = Inches(.65)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor.from_string(DARK)
for name, size, color in [('Title', 30, DARK), ('Heading 1', 21, DARK), ('Heading 2', 14, GREEN)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor.from_string(color); styles[name].font.bold = True

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None: node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def bookmark(paragraph, name, bookmark_id):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(bookmark_id)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark_id))
    paragraph._p.insert(0, start); paragraph._p.append(end)

def internal_link(paragraph, text, anchor):
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), anchor); link.set(qn('w:history'), '1')
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), GREEN)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(color); props.append(underline)
    node = OxmlElement('w:t'); node.text = text; run.append(props); run.append(node); link.append(run); paragraph._p.append(link)

def banner(text, color=PALE):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    c = t.cell(0,0); shade(c, color); set_cell_margins(c)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(GREEN if color == PALE else WHITE)

def footer_header():
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = h.add_run('ORVIX  •  MWONGOZO WA MOBILE'); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GREEN)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run('orvix.africa  •  info@orvix.africa  •  +255 655 746 552'); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)

def screenshot(name, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SHOTS/name), width=Inches(3.05))
    cp = doc.add_paragraph(caption); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic = True; cp.runs[0].font.size = Pt(8.5); cp.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

def steps(items):
    for item in items:
        p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(5); p.add_run(item)

def bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4); p.add_run(item)

def section(title, subtitle, image=None, caption=None, how=None, tips=None):
    doc.add_page_break(); heading = doc.add_heading(title, 1)
    number = title.split('.', 1)[0]; bookmark(heading, f'section_{number}', 100 + int(number))
    p = doc.add_paragraph(subtitle); p.runs[0].font.size = Pt(12); p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    if image: screenshot(image, caption or '')
    if how:
        doc.add_heading('Hatua kwa hatua', 2); steps(how)
    if tips:
        banner('KUMBUKA'); bullets(tips)

footer_header()

# Cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'public/icon-512.png'), width=Inches(1.25))
p = doc.add_paragraph('ORVIX', style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('MWONGOZO WA MTUMIAJI WA MOBILE'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.runs[0]; r.bold = True; r.font.size = Pt(17); r.font.color.rgb = RGBColor.from_string(GREEN)
p = doc.add_paragraph('Jinsi ya kutumia Orvix kwenye simu — hatua kwa hatua kwa Kiswahili rahisi.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
doc.add_paragraph('\n')
banner('TOLEO LA DEMO ACCOUNT  •  AUGUST 2026')
doc.add_paragraph('\n')
p = doc.add_paragraph('Picha zote ndani ya mwongozo huu zimetokana na demo account ya Orvix kwenye ukubwa wa simu. Data na majina ya bidhaa ni ya mafunzo.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(10); p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

doc.add_page_break(); doc.add_heading('Yaliyomo', 1)
toc = [
('1', 'Kuanza na kuingia'), ('2', 'Dashboard / Home'), ('3', 'Sales'), ('4', 'Sell / POS'), ('5', 'Stock'),
('6', 'More na features zote'), ('7', 'Buying / Purchases'), ('8', 'Expenses'), ('9', 'Reports'),
('10', 'Partners na Delivery'), ('11', 'Money na Planning'), ('12', 'Staff'), ('13', 'Settings, Sync na Subscription'), ('14', 'Msaada na mawasiliano')]
for n, title in toc:
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Inches(.55); t.columns[1].width = Inches(6.1)
    shade(t.cell(0,0), GREEN); shade(t.cell(0,1), LIGHT)
    p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(n); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    p=t.cell(0,1).paragraphs[0]; internal_link(p, title, f'section_{n}')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

section('1. Kuanza na kuingia', 'Tumia namba yako ya WhatsApp na Security PIN/Password.', '01-login.png', 'Screen ya kwanza ya Sign In kwenye simu.', [
    'Fungua Orvix kwenye browser ya simu au app iliyowekwa kwenye Home Screen.',
    'Chagua “Sign In”. Andika namba ya WhatsApp iliyosajiliwa.',
    'Bonyeza “Continue”, kisha andika Security PIN/Password.',
    'Bonyeza “Continue” tena. Orvix itafungua biashara na role yako.'
], ['Tumia namba iliyoalikwa na Admin ikiwa wewe ni staff.', 'Usimpe mtu mwingine PIN au password yako.'])
screenshot('01b-password.png', 'Hatua ya password baada ya namba kuthibitishwa.')

section('2. Dashboard / Home', 'Muhtasari wa biashara yako kwa haraka.', '02-dashboard.png', 'Dashboard halisi ya Orvix Mobile.', [
    'Angalia Total Orders, Today Sales, Purchases, Profit, Expenses na Dues Owed.',
    'Bonyeza “Open Sell Screen” kuanza kuuza mara moja.',
    'Tumia tabs za chini: Home, Sales, Sell, Stock na More.'
], ['Takwimu hubadilika kulingana na branch na muda uliochaguliwa.', 'Alama ya globe hubadilisha lugha.'])

section('3. Sales', 'Orodha ya mauzo, risiti na utafutaji wa transaction.', '03-sales.png', 'Sales screen kwenye simu.', [
    'Bonyeza “Sales” kwenye bottom navigation.',
    'Tumia Search kutafuta mauzo kwa jina, namba au kumbukumbu.',
    'Fungua sale kuona bidhaa, malipo na taarifa za mteja.',
    'Tumia action inayoruhusiwa na role yako, kwa mfano receipt au return.'
], ['Usifute au kurejesha sale bila kuthibitisha transaction.', 'Permissions za staff zinaweza kuficha baadhi ya actions.'])

section('4. Sell / POS', 'Screen ya kuuza bidhaa kwa retail au wholesale.', '04-pos.png', 'POS ya mobile yenye search, bidhaa na cart.', [
    'Bonyeza “Sell” katikati ya bottom navigation.',
    'Tafuta bidhaa kwa jina, code/barcode au tumia scanner.',
    'Chagua Retail Sales type au Wholesale Sales type.',
    'Bonyeza “+ ADD” kwa kila bidhaa; hakiki Cart Total.',
    'Bonyeza “Checkout”, chagua malipo, kisha thibitisha sale.'
], ['Thibitisha quantity na bei kabla ya Checkout.', 'Ikiwa bidhaa haionekani, hakiki branch stock na spelling ya search.'])

section('5. Stock', 'Kusimamia bidhaa, quantity, bei na stock movement.', '05-stock.png', 'Stock screen ya Orvix Mobile.', [
    'Bonyeza “Stock” kwenye bottom navigation.',
    'Tumia Search kutafuta bidhaa.',
    'Fungua bidhaa kuona quantity, buying price na selling price.',
    'Tumia Add/Edit/Adjust Stock kulingana na permission yako.'
], ['Mabadiliko ya stock yaingizwe kwenye branch sahihi.', 'Usitumie stock adjustment kuficha purchase au sale.'])

section('6. More na features zote', 'Hapa ndipo unapata menu zote za biashara.', '06-more.png', 'More menu: Buying, Expenses, Reports, Delivery, Partners, Money, Planning, Staff na System.', [
    'Bonyeza “More” upande wa chini kulia.',
    'Chagua feature unayotaka kwa kugusa icon yake.',
    'Tumia X juu kulia kufunga menu na kurudi kwenye screen ya awali.'
], ['Menu unazoona hutegemea role na permissions ulizopewa.', '“Sync — All data synced” inaonyesha hali ya data ya kifaa.'])

section('7. Buying / Purchases', 'Rekodi bidhaa zinazoingia kutoka kwa supplier.', None, None, [
    'Fungua More, kisha bonyeza “Buying”.',
    'Chagua supplier au ongeza taarifa zake.',
    'Ongeza bidhaa, quantity, buying price na gharama husika.',
    'Hakiki jumla na hali ya malipo, kisha Save/Complete Purchase.'
], ['Purchase iliyokamilika inaweza kuongeza stock.', 'Tumia invoice/reference ya supplier kwa ufuatiliaji.'])

section('8. Expenses', 'Rekodi matumizi ya biashara kwa category sahihi.', None, None, [
    'Fungua More, kisha “Expenses”.',
    'Bonyeza kuongeza expense mpya.',
    'Weka category, kiasi, tarehe, njia ya malipo na maelezo.',
    'Hifadhi na uhakikishe expense inaonekana kwenye list.'
], ['Ambatisha ushahidi wa malipo pale unapohitajika.', 'Usichanganye matumizi binafsi na ya biashara.'])

section('9. Reports', 'Chambua mauzo, stock, purchases, expenses na faida.', None, None, [
    'Fungua More, kisha “Reports”.',
    'Chagua report unayotaka.',
    'Weka muda: Today, Week, Month au custom range.',
    'Chuja kwa branch au category ikiwa option ipo.',
    'Soma summary na export/share ikiwa role yako inaruhusu.'
], ['Hakikisha date range kabla ya kulinganisha takwimu.', 'Report sahihi hutegemea data kuingizwa kwa usahihi.'])

section('10. Partners na Delivery', 'Simamia customers/suppliers na delivery orders.', None, None, [
    'Kutoka More, chagua “Partners” kuona customers na suppliers.',
    'Tafuta partner au ongeza mpya kwa jina na mawasiliano.',
    'Kwa delivery, chagua “Delivery”, fungua order na badilisha status kadri inavyotekelezwa.',
    'Thibitisha recipient, address na malipo kabla ya kukamilisha delivery.'
], ['Linda taarifa binafsi za customers.', 'Usiweke delivery Complete kabla bidhaa kufika.'])

section('11. Money na Planning', 'Fedha, akaunti, makadirio na mipango ya biashara.', None, None, [
    'Fungua “Money” kuona cash/bank/mobile money records zinazopatikana.',
    'Tumia “Planning” kwa forecasting na mipango inayoruhusiwa.',
    'Chagua muda na branch sahihi kabla ya kusoma takwimu.',
    'Linganisha entries na statement halisi ya malipo.'
], ['Reconciliation ifanywe mara kwa mara.', 'Forecast ni makadirio; si uthibitisho wa fedha halisi.'])

section('12. Staff', 'Alika staff, mpe role na dhibiti access.', None, None, [
    'Fungua More, kisha “Staff”.',
    'Bonyeza kuongeza/invite staff; weka jina, Gmail/mawasiliano na role.',
    'Chagua permissions zinazohitajika tu na branch husika.',
    'Tuma invite link. Staff aifungue, akamilishe usajili na aingie.',
    'Admin anaweza kusahihisha role au kuondoa access baadaye.'
], ['Usishirikishe invite link hadharani.', 'Staff anaweza kuingia kwenye kifaa tofauti kwa account ileile baada ya usajili.'])

section('13. Settings, Sync na Subscription', 'Mipangilio ya biashara, data sync na package.', None, None, [
    'Fungua More; chini ya System bonyeza “Settings”.',
    'Badilisha taarifa muhimu tu, kisha Save.',
    'Angalia “Sync” kuthibitisha data imeunganishwa.',
    'Fungua “Subscription” kuona package, expiry na njia ya kulipa.',
    'Tumia “Install Orvix App” kuongeza Orvix kwenye Home Screen.'
], ['Usifunge browser wakati mabadiliko muhimu yanahifadhiwa.', 'Renew package kabla haija-expire ili huduma isiingiliwe.'])

section('14. Msaada na mawasiliano', 'Timu ya Orvix ipo tayari kukusaidia.', None, None, [
    'Andaa screenshot ya tatizo na eleza ulikuwa kwenye menu gani.',
    'Taja aina ya simu, browser na hatua zilizosababisha tatizo.',
    'Wasiliana kupitia WhatsApp/Simu au Email hapa chini.'
], ['Usitume password/PIN kwenye ujumbe wa msaada.'])
banner('ORVIX SOLUTIONS  •  P.O. BOX 1259, MBEYA, TANZANIA', GREEN)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for line in ['WhatsApp / Simu: +255 655 746 552', 'Email: info@orvix.africa', 'Website: https://orvix.africa/']:
    r=p.add_run(line+'\n'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(DARK)

doc.save(OUT)
print(OUT)
